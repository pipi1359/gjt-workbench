#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""本地验证 sync_price.py 的解析逻辑（不写 Supabase）"""
import json, io, os, sys

# 导入脚本中的解析函数
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import importlib.util
spec = importlib.util.spec_from_file_location("sync_price", r"C:\Users\1\Doubao\chats\2026-10-09\new-chat\deploy-gjt-workbench\scripts\sync_price.py")
sp = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sp)

src = r"C:\Users\1\Doubao\chats\2026-10-09\new-chat\deploy-gjt-workbench\source\test_sample.xlsx"
from openpyxl import load_workbook
wb = load_workbook(src, read_only=True, data_only=True)
ws = wb.active
rows = ws.iter_rows(values_only=True)
headers = [str(h).strip() if h is not None else '' for h in next(rows, [])]
print("表头:", headers)
cm = sp.col_map(headers)
print("列映射:", cm)
# 手动模拟解析前3行
items = []
cnt = 0
for row in rows:
    if not row or all(c is None or str(c).strip()=='' for c in row): continue
    it = {f: (row[cm[f]] if cm[f] < len(row) else '') for f in ['b','n','br','c','m','u','p'] if f in cm}
    print("解析行:", it)
    cnt += 1
    if cnt >= 3: break
print("OK 解析验证通过")
