#!/usr/bin/env python3
"""
国补备案价 自动同步脚本
读取仓库 source/ 下的国补备案价 Excel 源文件，转换为 {generated,count,items}
写入 Supabase gjt_price 表 (id=1, payload)。

依赖: pip install openpyxl requests
GitHub Actions 环境变量:
  SUPABASE_URL   e.g. https://xxx.supabase.co
  SUPABASE_KEY   service_role key 或具备写权限的 key
  SOURCE_FILE    源文件名（默认国补备案价.xlsx，位于 source/ 下）
"""
import os, sys, json, io, datetime
import requests
from openpyxl import load_workbook

REPO = os.environ.get("GITHUB_WORKSPACE", os.getcwd())
SUPA_URL = os.environ.get("SUPABASE_URL", "https://xwwdksbxxhbqnlrvegus.supabase.co")
SUPA_KEY = os.environ.get("SUPABASE_KEY", "")
SOURCE = os.environ.get("SOURCE_FILE", "国补备案价.xlsx")
TABLE = "gjt_price"
ROW_ID = 1

def find_source():
    # 优先在仓库 source/ 下查找
    for base in [os.path.join(REPO, "source"), os.path.join(REPO, "source", "prices")]:
        p = os.path.join(base, SOURCE)
        if os.path.exists(p):
            return p
    # 否则在整个仓库里搜索
    for root, _dirs, files in os.walk(REPO):
        for f in files:
            if f.lower() == SOURCE.lower() or (os.path.splitext(f)[1].lower() in ('.xlsx','.xls') and '国补' in f):
                full = os.path.join(root, f)
                if 'source' in root or f == SOURCE:
                    return full
    raise FileNotFoundError(f"找不到源文件 {SOURCE} 于仓库 source/ 目录")

def col_map(headers):
    """把表头列名映射到标准字段 b/n/br/c/m/u/p"""
    aliases = {
        'b': ['条码','商品条码','barcode','条码号'],
        'n': ['名称','商品名称','name','品名'],
        'br': ['品牌','商品品牌','brand'],
        'c': ['分类','商品分类','category','类目'],
        'm': ['型号','商品型号','model','model_no'],
        'u': ['单位','unit'],
        'p': ['售价','价格','售价(元)','国补价','备案价','price','售价(元)'],
    }
    idx = {}
    norm = {k: (v.strip().lower() if v else '') for k, v in enumerate(headers)}
    for field, keys in aliases.items():
        for k, hv in norm.items():
            if any(ak in hv for ak in keys):
                idx[field] = k
                break
    return idx

def to_num(s):
    if s is None: return ''
    if isinstance(s, (int, float)): return s
    s = str(s).strip()
    s = s.replace(',', '').replace('￥','').replace('¥','').replace('元','').strip()
    if not s: return ''
    try:
        f = float(s)
        return f if f == int(f) else round(f, 2)
    except Exception:
        return s

def main():
    if not SUPA_KEY:
        print("缺少 SUPABASE_KEY，跳过同步", flush=True)
        sys.exit(1)
    src = find_source()
    print(f"读取源文件: {src}", flush=True)
    wb = load_workbook(src, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else '' for h in next(rows, [])]
    cm = col_map(headers)
    print(f"列映射: {cm}", flush=True)
    if not cm:
        print("警告：未能识别表头列，按固定顺序 b,n,br,c,m,u,p 解析首7列", flush=True)
    items = []
    seen = set()
    for row in rows:
        if not row or all(c is None or str(c).strip()=='' for c in row):
            continue
        if cm:
            it = {f: (row[cm[f]] if cm[f] < len(row) else '') for f in ['b','n','br','c','m','u','p'] if f in cm}
        else:
            it = {'b': row[0] if len(row)>0 else '', 'n': row[1] if len(row)>1 else '',
                  'br': row[2] if len(row)>2 else '', 'c': row[3] if len(row)>3 else '',
                  'm': row[4] if len(row)>4 else '', 'u': row[5] if len(row)>5 else '',
                  'p': row[6] if len(row)>6 else ''}
        b = str(it.get('b','')).strip()
        if not b or b.lower() in ('条码','商品条码','barcode'):  # 跳过重复表头
            continue
        if b in seen:
            continue
        seen.add(b)
        items.append({
            'b': b,
            'n': str(it.get('n','') or '').strip(),
            'br': str(it.get('br','') or '').strip(),
            'c': str(it.get('c','') or '').strip(),
            'm': str(it.get('m','') or '').strip(),
            'u': str(it.get('u','') or '').strip(),
            'p': to_num(it.get('p'))
        })
    wb.close()
    payload = {
        'generated': datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
        'count': len(items),
        'items': items
    }
    print(f"共解析 {len(items)} 条商品", flush=True)
    # 写入 Supabase (upsert)
    url = f"{SUPA_URL}/rest/v1/{TABLE}?on_conflict=id"
    body = [{'id': ROW_ID, 'payload': payload, 'updated_at': datetime.datetime.utcnow().isoformat() + 'Z'}]
    headers = {
        'apikey': SUPA_KEY,
        'Authorization': 'Bearer ' + SUPA_KEY,
        'Content-Type': 'application/json',
        'Prefer': 'resolution=merge-duplicates,return=minimal'
    }
    r = requests.post(url, json=body, headers=headers, timeout=120)
    print(f"Supabase 写入状态: {r.status_code}", flush=True)
    if r.status_code not in (200, 201, 204):
        print("写入失败: " + r.text[:500], flush=True)
        sys.exit(1)
    print(f"同步完成: {len(items)} 条 -> {TABLE}", flush=True)

if __name__ == '__main__':
    main()
